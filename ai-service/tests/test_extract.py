import asyncio
from io import BytesIO

import httpx
import pymupdf
from docx import Document

from app.core.config import get_settings
from app.main import app
from app.services.text_normalization import normalize_extracted_text

DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def post_file(
    content: bytes,
    content_type: str,
    filename: str = "sample",
) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(
            transport=transport,
            base_url="http://test",
        ) as client:
            return await client.post(
                "/extract",
                files={"file": (filename, content, content_type)},
            )

    return asyncio.run(send())


def make_pdf(pages: list[str | None], *, image_only: bool = False) -> bytes:
    document = pymupdf.open()
    for page_text in pages:
        page = document.new_page()
        if page_text:
            page.insert_text((72, 72), page_text)
        if image_only:
            pixmap = pymupdf.Pixmap(
                pymupdf.csRGB,
                pymupdf.IRect(0, 0, 4, 4),
                False,
            )
            pixmap.clear_with(255)
            page.insert_image(page.rect, stream=pixmap.tobytes("png"))
            page.draw_rect(
                pymupdf.Rect(72, 72, 210, 210),
                color=(0, 0, 0),
                fill=(0, 0, 0),
            )

    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left cell"
    table.cell(0, 1).text = "right cell"
    document.add_paragraph("Last paragraph")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_extract_valid_pdf_reports_text_and_page_count() -> None:
    response = post_file(
        make_pdf(["A sufficiently long sentence for PDF extraction testing."]),
        "application/pdf",
        "sample.pdf",
    )

    assert response.status_code == 200
    result = response.json()
    assert result["status"] == "success"
    assert result["fileType"] == "pdf"
    assert result["text"].startswith("A sufficiently long sentence")
    assert result["pageCount"] == 1
    assert result["wordCount"] == len(result["text"].split())
    assert result["characterCount"] == len(result["text"])
    assert result["requiresOcr"] is False
    assert result["warnings"] == []


def test_extract_rejects_corrupt_pdf_without_parser_details() -> None:
    response = post_file(b"%PDF-1.7\nnot a real PDF", "application/pdf", "bad.pdf")

    assert response.status_code == 422
    assert response.json() == {"detail": "The PDF is invalid or could not be read."}
    assert "traceback" not in response.text.lower()


def test_extract_rejects_pdf_mime_magic_mismatch() -> None:
    response = post_file(b"not a PDF", "application/pdf", "bad.pdf")

    assert response.status_code == 400
    assert "does not match" in response.json()["detail"]


def test_extract_marks_image_only_pdf_as_requiring_ocr() -> None:
    response = post_file(
        make_pdf([None], image_only=True),
        "application/pdf",
        "scan.pdf",
    )

    assert response.status_code == 200
    result = response.json()
    assert result["pageCount"] == 1
    assert result["text"] == ""
    assert result["requiresOcr"] is True
    assert result["warnings"] == [
        "PDF has limited extractable text; OCR may be required."
    ]


def test_pdf_ocr_rule_uses_configured_low_text_page_ratio(monkeypatch) -> None:
    monkeypatch.setenv("OCR_MIN_CHARS_PER_PAGE", "20")
    monkeypatch.setenv("OCR_MIN_LOW_TEXT_PAGE_RATIO", "0.75")
    get_settings.cache_clear()

    try:
        response = post_file(
            make_pdf(["Sufficient text for average density. " * 8, None]),
            "application/pdf",
            "two-pages.pdf",
        )
    finally:
        get_settings.cache_clear()

    assert response.status_code == 200
    result = response.json()
    assert result["pageCount"] == 2
    assert result["requiresOcr"] is False
    assert result["warnings"] == []


def test_extract_valid_docx_preserves_paragraph_and_table_order() -> None:
    response = post_file(make_docx(), DOCX_MIME_TYPE, "sample.docx")

    assert response.status_code == 200
    result = response.json()
    assert result["fileType"] == "docx"
    assert result["text"] == "First paragraph\nleft cell\tright cell\nLast paragraph"
    assert result["pageCount"] is None
    assert result["wordCount"] == 8
    assert result["characterCount"] == len(result["text"])
    assert result["requiresOcr"] is False
    assert result["warnings"] == []


def test_extract_rejects_corrupt_docx() -> None:
    response = post_file(b"not a DOCX zip", DOCX_MIME_TYPE, "bad.docx")

    assert response.status_code == 422
    assert response.json() == {
        "detail": "The DOCX file is invalid or could not be read."
    }


def test_extract_txt_decodes_utf8_bom_and_normalizes_line_breaks() -> None:
    response = post_file(
        b"\xef\xbb\xbf  A  \r\n\r\n\r\n\r\nB ",
        "text/plain; charset=utf-8",
        "sample.txt",
    )

    assert response.status_code == 200
    result = response.json()
    assert result["fileType"] == "txt"
    assert result["text"] == "  A  \n\n\nB "
    assert result["pageCount"] is None
    assert result["wordCount"] == 2
    assert result["characterCount"] == len(result["text"])
    assert result["requiresOcr"] is False
    assert result["warnings"] == []


def test_extract_txt_uses_windows_1252_fallback() -> None:
    response = post_file(b"caf\xe9", "text/plain", "sample.txt")

    assert response.status_code == 200
    result = response.json()
    assert result["text"] == "café"
    assert result["warnings"] == ["Text was decoded using the Windows-1252 fallback."]
    assert result["characterCount"] == 4
    assert result["wordCount"] == 1


def test_extract_txt_uses_latin_1_fallback_when_cp1252_cannot_decode() -> None:
    content = b"Fallback text with a Latin-1 byte: \x81 retained."
    response = post_file(content, "text/plain", "sample.txt")

    assert response.status_code == 200
    result = response.json()
    assert result["text"] == content.decode("latin-1")
    assert result["warnings"] == ["Text was decoded using the Latin-1 fallback."]


def test_extract_rejects_binary_txt_with_nul_or_controls() -> None:
    for content in (
        b"plain\x00binary",
        b"\x01\x02\x03",
        b"\x81\x82",
        b"%PDF-1.7\nnot plain text",
    ):
        response = post_file(content, "text/plain", "binary.txt")

        assert response.status_code == 422
        detail = response.json()["detail"]
        assert "binary" in detail.lower() or "NUL" in detail


def test_extract_rejects_unsupported_media_type() -> None:
    response = post_file(b"text", "application/octet-stream", "sample.bin")

    assert response.status_code == 415


def test_normalization_removes_nuls_only_for_safe_extracted_formats() -> None:
    assert normalize_extracted_text("a\x00b", remove_nuls=True) == "ab"
    assert normalize_extracted_text("a\x00b", remove_nuls=False) == "a\x00b"


def test_file_contract_response_has_only_documented_fields() -> None:
    response = post_file(b"hello world", "text/plain", "sample.txt")

    assert set(response.json()) == {
        "status",
        "fileType",
        "text",
        "pageCount",
        "wordCount",
        "characterCount",
        "requiresOcr",
        "warnings",
        "pages",
    }


def test_extract_enforces_configured_max_upload_size(monkeypatch) -> None:
    monkeypatch.setenv("MAX_UPLOAD_MB", "1")
    get_settings.cache_clear()

    try:
        response = post_file(
            b"x" * (1024 * 1024 + 1),
            "text/plain",
            "large.txt",
        )
    finally:
        get_settings.cache_clear()

    assert response.status_code == 413
    assert response.json() == {"detail": "File exceeds the 1 MB upload limit."}
